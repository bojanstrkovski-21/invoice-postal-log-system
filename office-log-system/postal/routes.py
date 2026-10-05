import csv
import io

from auth import admin_required, login_required
from db import get_db
from flask import Blueprint, current_app, jsonify, request

postal_bp = Blueprint('postal', __name__, url_prefix='/api/postal')

PACKAGE_TYPES = (
    'обична',
    'препорачана',
    'препорачана со повратница',
    'брза пошта',
    'брза пошта препорачана со повратница',
)


def _clean(value):
    if value is None:
        return ''
    return str(value).strip()


def _optional(value):
    return _clean(value) or None


def _partner_id(value):
    if value in (None, '', 0, '0'):
        return None
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def _filters(args):
    clauses = ['1=1']
    params = []
    search = _clean(args.get('search'))
    year = _clean(args.get('year'))
    date_from = _clean(args.get('date_from'))
    date_to = _clean(args.get('date_to'))
    if search:
        clauses.append(
            '(p.recipient_name LIKE ? OR IFNULL(p.tax_number, \'\') LIKE ? OR p.ref_number LIKE ?)'
        )
        like = f'%{search}%'
        params.extend([like, like, like])
    if year:
        clauses.append('p.year = ?')
        params.append(int(year))
    if date_from:
        clauses.append('p.send_date >= ?')
        params.append(date_from)
    if date_to:
        clauses.append('p.send_date <= ?')
        params.append(date_to)
    return ' AND '.join(clauses), params


def _parse_package(data):
    send_date = _clean(data.get('send_date'))
    ref_number = _clean(data.get('ref_number'))
    recipient_name = _clean(data.get('recipient_name'))
    package_type = _clean(data.get('package_type')) or PACKAGE_TYPES[0]
    if not send_date or len(send_date) < 4 or not send_date[:4].isdigit():
        return None, 'Date is required'
    if not ref_number or not recipient_name:
        return None, 'Reference number and recipient are required'
    if package_type not in PACKAGE_TYPES:
        return None, 'Invalid package type'
    return (
        ref_number,
        send_date,
        _partner_id(data.get('partner_id')),
        recipient_name,
        _optional(data.get('recipient_address')),
        _optional(data.get('city')),
        _optional(data.get('tax_number')),
        package_type,
        int(send_date[:4]),
        _optional(data.get('notes')),
    ), None


def _list_query(where, params, order):
    db = get_db()
    sql = f'''
        SELECT p.*, partners.name AS partner_name
        FROM packages p
        LEFT JOIN partners ON partners.id = p.partner_id
        WHERE {where}
        ORDER BY {order}
    '''
    return [dict(row) for row in db.execute(sql, params).fetchall()]


@postal_bp.route('/status')
@login_required
def status():
    return jsonify({'module': 'postal', 'ready': True})


@postal_bp.route('/packages', methods=['GET'])
@login_required
def get_packages():
    where, params = _filters(request.args)
    rows = _list_query(where, params, 'p.send_date DESC, p.id DESC')
    return jsonify(rows)


@postal_bp.route('/packages', methods=['POST'])
@login_required
def create_package():
    data = request.get_json(silent=True) or {}
    values, error = _parse_package(data)
    if error or values is None:
        return jsonify({'error': error or 'Invalid package'}), 400
    db = get_db()
    cur = db.execute(
        '''INSERT INTO packages (
               ref_number, send_date, partner_id, recipient_name, recipient_address,
               city, tax_number, package_type, year, notes
           ) VALUES (?,?,?,?,?,?,?,?,?,?)''',
        values
    )
    db.commit()
    row = _list_query('p.id = ?', [cur.lastrowid], 'p.id')[0]
    return jsonify(row), 201


@postal_bp.route('/packages/<int:pkg_id>', methods=['PUT'])
@admin_required
def update_package(pkg_id):
    data = request.get_json(silent=True) or {}
    values, error = _parse_package(data)
    if error or values is None:
        return jsonify({'error': error or 'Invalid package'}), 400
    db = get_db()
    existing = db.execute('SELECT id FROM packages WHERE id = ?', (pkg_id,)).fetchone()
    if not existing:
        return jsonify({'error': 'Not found'}), 404
    db.execute(
        '''UPDATE packages SET
               ref_number=?, send_date=?, partner_id=?, recipient_name=?,
               recipient_address=?, city=?, tax_number=?, package_type=?,
               year=?, notes=?
           WHERE id=?''',
        (*values, pkg_id)
    )
    db.commit()
    return jsonify(_list_query('p.id = ?', [pkg_id], 'p.id')[0])


@postal_bp.route('/packages/<int:pkg_id>', methods=['DELETE'])
@admin_required
def delete_package(pkg_id):
    db = get_db()
    db.execute('DELETE FROM packages WHERE id = ?', (pkg_id,))
    db.commit()
    return jsonify({'ok': True})


@postal_bp.route('/years')
@login_required
def get_years():
    db = get_db()
    rows = db.execute('SELECT DISTINCT year FROM packages ORDER BY year DESC').fetchall()
    return jsonify([row['year'] for row in rows])


@postal_bp.route('/export')
@login_required
def export_packages():
    where, params = _filters(request.args)
    rows = _list_query(where, params, 'p.send_date ASC, p.id ASC')
    headers = ['Датум', 'Деловоден број', 'Примач', 'Даночен број', 'Адреса', 'Град', 'Вид пратка', 'Забелешка']
    fmt = request.args.get('format', 'csv')

    def cells(row):
        return [
            row['send_date'], row['ref_number'], row['recipient_name'],
            row['tax_number'] or '', row['recipient_address'] or '', row['city'] or '',
            row['package_type'], row['notes'] or ''
        ]

    if fmt == 'xlsx':
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill
        except ImportError:
            return jsonify({'error': 'openpyxl not installed'}), 500
        wb = openpyxl.Workbook()
        ws = wb.active
        if ws is None:
            ws = wb.create_sheet('Пратки')
        ws.title = 'Пратки'
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='2563EB')
        for row in rows:
            ws.append(cells(row))
        widths = [12, 18, 28, 18, 28, 16, 36, 24]
        for index, width in enumerate(widths, 1):
            ws.column_dimensions[chr(64 + index)].width = width
        buf = io.BytesIO()
        wb.save(buf)
        return _download(
            buf.getvalue(),
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            'pratki.xlsx'
        )
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(headers)
    for row in rows:
        writer.writerow(cells(row))
    return _download(buf.getvalue().encode('utf-8-sig'), 'text/csv; charset=utf-8', 'pratki.csv')


def _download(content, mimetype, filename):
    return current_app.response_class(
        content,
        mimetype=mimetype,
        headers={'Content-Disposition': f'attachment; filename={filename}'}
    )
