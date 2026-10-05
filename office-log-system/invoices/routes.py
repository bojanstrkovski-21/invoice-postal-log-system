import csv
import io
import sqlite3
from datetime import date, datetime, timezone

from auth import admin_required, login_required
from db import get_db
from flask import Blueprint, current_app, jsonify, request
from nbrm import NbrmError, lookup_rate

invoices_bp = Blueprint('invoices', __name__, url_prefix='/api/invoices')

INTERNAL_PREFIX = '05-'
PAYMENT_STATUSES = ('pending', 'paid', 'cancelled')
STATUS_LABELS = {
    'pending': 'неплатена',
    'paid': 'платена',
    'cancelled': 'откажана',
}


def _clean(value):
    if value is None:
        return ''
    return str(value).strip()


def _optional(value):
    return _clean(value) or None


def _partner_id(value):
    if value in (None, '', 0, '0'):
        return None
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def _date(value, required=False):
    text = _clean(value)
    if not text:
        if required:
            return None, 'Date is required'
        return None, None
    try:
        date.fromisoformat(text)
    except ValueError:
        return None, 'Invalid date'
    return text, None


def _amount(value):
    if value is None or value == '':
        return None, None
    if isinstance(value, bool):
        return None, 'Invalid amount'
    if isinstance(value, (int, float)):
        return float(value), None
    text = _clean(value).replace(' ', '').replace('\u00a0', '')
    if not text:
        return None, None
    if ',' in text and '.' in text:
        if text.rfind(',') > text.rfind('.'):
            text = text.replace('.', '').replace(',', '.')
        else:
            text = text.replace(',', '')
    elif ',' in text:
        text = text.replace(',', '.')
    elif text.count('.') > 1:
        text = text.replace('.', '')
    elif text.count('.') == 1:
        left, right = text.split('.')
        if len(right) == 3 and left.isdigit() and right.isdigit():
            text = left + right
    try:
        return float(text), None
    except ValueError:
        return None, 'Invalid amount'


ACCOUNTS = {
    '787': '787 — сметка за самофинансирачки активности',
    '903': '903 — сметка за наменска дотација',
    '785': '785 — донаторски сметки',
}


def format_amount(value):
    if value is None:
        return ''
    negative = float(value) < 0
    whole, frac = f'{abs(float(value)):.2f}'.split('.')
    grouped = ''
    for index, char in enumerate(reversed(whole)):
        if index and index % 3 == 0:
            grouped = '.' + grouped
        grouped = char + grouped
    return ('-' if negative else '') + grouped + ',' + frac


def format_rate(value):
    if value is None:
        return ''
    whole, frac = f'{float(value):.4f}'.split('.')
    return whole + ',' + frac


def account_label(code):
    if not code:
        return ''
    return ACCOUNTS.get(code, code)


def next_internal_number(db, year):
    rows = db.execute(
        'SELECT internal_number FROM invoices WHERE year = ?',
        (year,)
    ).fetchall()
    highest = 0
    for row in rows:
        number = row['internal_number'] or ''
        if not number.startswith(INTERNAL_PREFIX):
            continue
        suffix = number[len(INTERNAL_PREFIX):]
        if suffix.isdigit():
            highest = max(highest, int(suffix))
    return f'{INTERNAL_PREFIX}{highest + 1}'


def _filters(args):
    clauses = ['1=1']
    params = []
    search = _clean(args.get('search'))
    year = _clean(args.get('year'))
    date_from = _clean(args.get('date_from'))
    date_to = _clean(args.get('date_to'))
    status = _clean(args.get('status'))
    if search:
        clauses.append(
            '''(IFNULL(i.supplier_name, '') LIKE ?
                OR IFNULL(partners.name, '') LIKE ?
                OR IFNULL(i.tax_number, '') LIKE ?
                OR IFNULL(i.invoice_number, '') LIKE ?
                OR i.internal_number LIKE ?
                OR IFNULL(i.bank_account, '') LIKE ?)'''
        )
        like = f'%{search}%'
        params.extend([like, like, like, like, like, like])
    if year.isdigit():
        clauses.append('i.year = ?')
        params.append(int(year))
    if date_from:
        clauses.append('i.received_date >= ?')
        params.append(date_from)
    if date_to:
        clauses.append('i.received_date <= ?')
        params.append(date_to)
    if status in PAYMENT_STATUSES:
        clauses.append('i.payment_status = ?')
        params.append(status)
    return ' AND '.join(clauses), params


def _parse_invoice(data):
    received_date, error = _date(data.get('received_date'), required=True)
    if error:
        return None, error
    invoice_date, error = _date(data.get('invoice_date'))
    if error:
        return None, error
    due_date, error = _date(data.get('due_date'))
    if error or not received_date:
        return None, error or 'Date is required'
    supplier_name = _clean(data.get('supplier_name'))
    if not supplier_name:
        return None, 'Supplier is required'
    amount, error = _amount(data.get('amount'))
    if error:
        return None, error
    status = _clean(data.get('payment_status')) or 'pending'
    if status not in PAYMENT_STATUSES:
        return None, 'Invalid payment status'
    currency = (_clean(data.get('currency')) or 'MKD').upper()
    if currency not in ('MKD', 'EUR', 'USD'):
        return None, 'Invalid currency'
    internal = _clean(data.get('internal_number'))
    if len(internal) > 40:
        return None, 'Invalid internal number'
    return {
        'internal_number': internal,
        'invoice_number': _optional(data.get('invoice_number')),
        'invoice_date': invoice_date,
        'received_date': received_date,
        'due_date': due_date,
        'partner_id': _partner_id(data.get('partner_id')),
        'supplier_name': supplier_name,
        'tax_number': _optional(data.get('tax_number')),
        'bank_account': _optional(data.get('bank_account')),
        'amount': amount,
        'currency': currency,
        'exchange_rate': None,
        'rate_date': None,
        'amount_mkd': None,
        'rate_warning': None,
        'payment_status': status,
        'notes': _optional(data.get('notes')),
        'year': int(received_date[:4]),
    }, None


def _values(parsed):
    return (
        parsed['internal_number'],
        parsed['invoice_number'],
        parsed['invoice_date'],
        parsed['received_date'],
        parsed['due_date'],
        parsed['partner_id'],
        parsed['supplier_name'],
        parsed['tax_number'],
        parsed['bank_account'],
        parsed['amount'],
        parsed['currency'],
        parsed['exchange_rate'],
        parsed['rate_date'],
        parsed['amount_mkd'],
        parsed['payment_status'],
        parsed['notes'],
        parsed['year'],
    )


def _list_query(where, params, order):
    db = get_db()
    sql = f'''
        SELECT i.*, partners.name AS partner_name
        FROM invoices i
        LEFT JOIN partners ON partners.id = i.partner_id
        WHERE {where}
        ORDER BY {order}
    '''
    rows = []
    for row in db.execute(sql, params).fetchall():
        item = dict(row)
        if not item.get('supplier_name'):
            item['supplier_name'] = item.get('partner_name') or ''
        rows.append(item)
    return rows


def _duplicate():
    return jsonify({'error': 'Internal number already exists for this year'}), 409


@invoices_bp.route('/status')
@login_required
def status():
    return jsonify({'module': 'invoices', 'ready': True})


@invoices_bp.route('/next')
@login_required
def next_number():
    received, error = _date(request.args.get('received_date'))
    if error:
        return jsonify({'error': error}), 400
    year_arg = _clean(request.args.get('year'))
    if received:
        year = int(received[:4])
    elif year_arg.isdigit():
        year = int(year_arg)
    else:
        year = datetime.now(timezone.utc).astimezone().year
    return jsonify({
        'year': year,
        'internal_number': next_internal_number(get_db(), year),
    })


def apply_exchange(parsed):
    if parsed['currency'] == 'MKD':
        parsed['exchange_rate'] = 1
        parsed['rate_date'] = parsed['received_date']
        parsed['amount_mkd'] = parsed['amount']
        parsed['rate_warning'] = None
        return parsed
    try:
        rate = lookup_rate(parsed['currency'], parsed['received_date'])
    except NbrmError as exc:
        parsed['exchange_rate'] = None
        parsed['rate_date'] = None
        parsed['amount_mkd'] = None
        parsed['rate_warning'] = str(exc)
        return parsed
    parsed['exchange_rate'] = rate['rate']
    parsed['rate_date'] = rate['list_date']
    if parsed['amount'] is None:
        parsed['amount_mkd'] = None
    else:
        parsed['amount_mkd'] = round(parsed['amount'] * rate['rate'], 2)
    parsed['rate_warning'] = None
    return parsed


def _with_warning(row, warning):
    if warning:
        row = dict(row)
        row['rate_warning'] = warning
    return row


@invoices_bp.route('/rate')
@login_required
def exchange_rate():
    currency = _clean(request.args.get('currency')).upper() or 'MKD'
    day, error = _date(request.args.get('date'), required=True)
    if error:
        return jsonify({'error': error}), 400
    force = request.args.get('refresh') == '1'
    try:
        rate = lookup_rate(currency, day, force=force)
    except NbrmError as exc:
        return jsonify({'error': str(exc)}), 502
    return jsonify(rate)


@invoices_bp.route('', methods=['GET'])
@login_required
def get_invoices():
    where, params = _filters(request.args)
    rows = _list_query(where, params, 'i.received_date DESC, i.id DESC')
    return jsonify(rows)


@invoices_bp.route('', methods=['POST'])
@login_required
def create_invoice():
    data = request.get_json(silent=True) or {}
    parsed, error = _parse_invoice(data)
    if error or parsed is None:
        return jsonify({'error': error or 'Invalid invoice'}), 400
    apply_exchange(parsed)
    db = get_db()
    db.execute('BEGIN IMMEDIATE')
    try:
        if not parsed['internal_number']:
            parsed['internal_number'] = next_internal_number(db, parsed['year'])
        cur = db.execute(
            '''INSERT INTO invoices (
                   internal_number, invoice_number, invoice_date, received_date, due_date,
                   partner_id, supplier_name, tax_number, bank_account, amount, currency,
                   exchange_rate, rate_date, amount_mkd, payment_status, notes, year
               ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)''',
            _values(parsed)
        )
        db.commit()
    except sqlite3.IntegrityError:
        db.rollback()
        return _duplicate()
    row = _list_query('i.id = ?', [cur.lastrowid], 'i.id')[0]
    return jsonify(_with_warning(row, parsed['rate_warning'])), 201


@invoices_bp.route('/<int:invoice_id>', methods=['PUT'])
@admin_required
def update_invoice(invoice_id):
    data = request.get_json(silent=True) or {}
    parsed, error = _parse_invoice(data)
    if error or parsed is None:
        return jsonify({'error': error or 'Invalid invoice'}), 400
    if not parsed['internal_number']:
        return jsonify({'error': 'Internal number is required'}), 400
    apply_exchange(parsed)
    db = get_db()
    existing = db.execute('SELECT id FROM invoices WHERE id = ?', (invoice_id,)).fetchone()
    if not existing:
        return jsonify({'error': 'Not found'}), 404
    try:
        db.execute(
            '''UPDATE invoices SET
                   internal_number=?, invoice_number=?, invoice_date=?, received_date=?,
                   due_date=?, partner_id=?, supplier_name=?, tax_number=?, bank_account=?,
                   amount=?, currency=?, exchange_rate=?, rate_date=?, amount_mkd=?,
                   payment_status=?, notes=?, year=?
               WHERE id=?''',
            (*_values(parsed), invoice_id)
        )
        db.commit()
    except sqlite3.IntegrityError:
        db.rollback()
        return _duplicate()
    row = _list_query('i.id = ?', [invoice_id], 'i.id')[0]
    return jsonify(_with_warning(row, parsed['rate_warning']))


@invoices_bp.route('/<int:invoice_id>', methods=['DELETE'])
@admin_required
def delete_invoice(invoice_id):
    db = get_db()
    db.execute('DELETE FROM invoices WHERE id = ?', (invoice_id,))
    db.commit()
    return jsonify({'ok': True})


@invoices_bp.route('/years')
@login_required
def get_years():
    db = get_db()
    rows = db.execute('SELECT DISTINCT year FROM invoices ORDER BY year DESC').fetchall()
    return jsonify([row['year'] for row in rows])


@invoices_bp.route('/export')
@login_required
def export_invoices():
    where, params = _filters(request.args)
    rows = _list_query(where, params, 'i.received_date ASC, i.id ASC')
    headers = [
        'Внатрешен број', 'Број на фактура', 'Датум на фактура', 'Назив',
        'Даночен број', 'Износ', 'Валута', 'Курс НБРМ', 'Износ во денари',
        'Сметка', 'Датум на прием',
        'Рок на плаќање', 'Статус', 'Забелешка',
    ]
    fmt = request.args.get('format', 'csv')

    def cells(row):
        return [
            row['internal_number'],
            row['invoice_number'] or '',
            row['invoice_date'] or '',
            row['supplier_name'] or '',
            row['tax_number'] or '',
            format_amount(row['amount']),
            row['currency'] or 'MKD',
            format_rate(row.get('exchange_rate')),
            format_amount(row.get('amount_mkd')),
            account_label(row['bank_account']),
            row['received_date'],
            row['due_date'] or '',
            STATUS_LABELS.get(row['payment_status'], row['payment_status']),
            row['notes'] or '',
        ]

    if fmt == 'xlsx':
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill
        except ImportError:
            return jsonify({'error': 'openpyxl not installed'}), 500
        wb = openpyxl.Workbook()
        ws = wb.active
        if ws is None:
            ws = wb.create_sheet('Фактури')
        ws.title = 'Фактури'
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='2563EB')
        for row in rows:
            ws.append(cells(row))
        widths = [18, 18, 16, 28, 18, 14, 10, 14, 16, 42, 16, 16, 14, 24]
        for index, width in enumerate(widths, 1):
            ws.column_dimensions[chr(64 + index)].width = width
        buf = io.BytesIO()
        wb.save(buf)
        return _download(
            buf.getvalue(),
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            'fakturi.xlsx'
        )
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(headers)
    for row in rows:
        writer.writerow(cells(row))
    return _download(buf.getvalue().encode('utf-8-sig'), 'text/csv; charset=utf-8', 'fakturi.csv')


def _download(content, mimetype, filename):
    return current_app.response_class(
        content,
        mimetype=mimetype,
        headers={'Content-Disposition': f'attachment; filename={filename}'}
    )
